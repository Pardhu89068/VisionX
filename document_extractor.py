"""
VisionX document text extraction engine.

Supported analysis:
PDF, DOCX, XLSX, XLS, PPTX, TXT, CSV, PNG, JPG, JPEG.

Other file types can still be stored safely in the Privacy Vault.
"""

import csv
import os
from pathlib import Path


SUPPORTED_ANALYSIS_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".xls",
    ".pptx",
    ".txt",
    ".csv",
    ".png",
    ".jpg",
    ".jpeg",
}


DISPLAY_NAMES = {
    ".pdf": "PDF",
    ".docx": "Word Document",
    ".xlsx": "Excel Workbook",
    ".xls": "Excel Spreadsheet",
    ".pptx": "PowerPoint Presentation",
    ".txt": "Text File",
    ".csv": "CSV File",
    ".png": "PNG Image",
    ".jpg": "JPG Image",
    ".jpeg": "JPEG Image",
    ".doc": "Legacy Word Document",
}


def get_file_extension(filename):
    return Path(filename or "").suffix.lower()


def get_file_type_label(filename):
    extension = get_file_extension(filename)

    return DISPLAY_NAMES.get(
        extension,
        extension.upper().replace(".", "") or "Unknown"
    )


def is_analysis_supported(filename):
    return (
        get_file_extension(filename)
        in SUPPORTED_ANALYSIS_EXTENSIONS
    )


def _read_text_file(file_path):
    encodings = (
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1"
    )

    for encoding in encodings:
        try:
            with open(
                file_path,
                "r",
                encoding=encoding,
                errors="strict"
            ) as file:
                return file.read()

        except UnicodeDecodeError:
            continue

    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="replace"
    ) as file:
        return file.read()


def _extract_pdf(file_path):
    from pypdf import PdfReader

    reader = PdfReader(file_path)

    pages = []

    for page in reader.pages:
        pages.append(
            page.extract_text() or ""
        )

    return "\n".join(pages).strip()


def _extract_docx(file_path):
    from docx import Document

    document = Document(file_path)

    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            if cells:
                parts.append(
                    " | ".join(cells)
                )

    return "\n".join(parts).strip()


def _extract_xlsx(file_path):
    from openpyxl import load_workbook

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )

    parts = []

    try:
        for worksheet in workbook.worksheets:

            parts.append(
                f"Sheet: {worksheet.title}"
            )

            for row in worksheet.iter_rows(
                values_only=True
            ):
                values = [
                    str(value).strip()
                    for value in row
                    if (
                        value is not None
                        and str(value).strip()
                    )
                ]

                if values:
                    parts.append(
                        " | ".join(values)
                    )

    finally:
        workbook.close()

    return "\n".join(parts).strip()


def _extract_xls(file_path):
    import xlrd

    workbook = xlrd.open_workbook(
        file_path,
        on_demand=True
    )

    parts = []

    try:
        for sheet_name in workbook.sheet_names():

            worksheet = workbook.sheet_by_name(
                sheet_name
            )

            parts.append(
                f"Sheet: {sheet_name}"
            )

            for row_index in range(
                worksheet.nrows
            ):
                values = []

                for column_index in range(
                    worksheet.ncols
                ):
                    value = worksheet.cell_value(
                        row_index,
                        column_index
                    )

                    if (
                        value is not None
                        and str(value).strip()
                    ):
                        values.append(
                            str(value).strip()
                        )

                if values:
                    parts.append(
                        " | ".join(values)
                    )

    finally:
        workbook.release_resources()

    return "\n".join(parts).strip()


def _extract_pptx(file_path):
    from pptx import Presentation

    presentation = Presentation(file_path)

    parts = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        parts.append(
            f"Slide {slide_number}"
        )

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    parts.append(text)

    return "\n".join(parts).strip()


def _extract_csv(file_path):
    parts = []

    with open(
        file_path,
        "r",
        encoding="utf-8-sig",
        errors="replace",
        newline=""
    ) as file:

        reader = csv.reader(file)

        for row in reader:

            values = [
                str(value).strip()
                for value in row
                if str(value).strip()
            ]

            if values:
                parts.append(
                    " | ".join(values)
                )

    return "\n".join(parts).strip()


def _extract_image(file_path):

    try:
        from PIL import Image
        import pytesseract

    except ImportError:
        return (
            "",
            "Image OCR is unavailable because "
            "the required Python packages are not installed."
        )

    try:
        image = Image.open(file_path)

        text = pytesseract.image_to_string(
            image
        )

        if text.strip():
            return (
                text.strip(),
                "Text extracted using local OCR."
            )

        return (
            "",
            "No readable text was detected in this image."
        )

    except Exception as error:

        return (
            "",
            "Image OCR is unavailable. "
            "Install the local Tesseract OCR engine "
            "to analyse images. "
            f"Technical detail: {error}"
        )


def extract_text(file_path):
    """
    Extract readable text from a supported file.

    Returns:

    {
        "supported": bool,
        "text": str,
        "message": str,
        "file_type": str
    }
    """

    filename = os.path.basename(
        file_path
    )

    extension = get_file_extension(
        filename
    )

    file_type = get_file_type_label(
        filename
    )

    # -------------------------------------------------
    # Unsupported file
    # -------------------------------------------------

    if extension not in SUPPORTED_ANALYSIS_EXTENSIONS:

        if extension == ".doc":

            return {
                "supported": False,
                "text": "",
                "message": (
                    "Legacy .DOC files can be stored "
                    "safely, but direct analysis is "
                    "unavailable without an external "
                    "converter."
                ),
                "file_type": file_type,
            }

        return {
            "supported": False,
            "text": "",
            "message": (
                f"{file_type} files can be stored "
                "safely, but content analysis is "
                "not currently supported."
            ),
            "file_type": file_type,
        }

    # -------------------------------------------------
    # Extract
    # -------------------------------------------------

    try:

        if extension == ".pdf":

            text = _extract_pdf(
                file_path
            )

            message = (
                "Text extracted from PDF."
            )

        elif extension == ".docx":

            text = _extract_docx(
                file_path
            )

            message = (
                "Text extracted from Word document."
            )

        elif extension == ".xlsx":

            text = _extract_xlsx(
                file_path
            )

            message = (
                "Data extracted from Excel workbook."
            )

        elif extension == ".xls":

            text = _extract_xls(
                file_path
            )

            message = (
                "Data extracted from Excel spreadsheet."
            )

        elif extension == ".pptx":

            text = _extract_pptx(
                file_path
            )

            message = (
                "Text extracted from PowerPoint presentation."
            )

        elif extension == ".txt":

            text = _read_text_file(
                file_path
            )

            message = (
                "Text extracted from text file."
            )

        elif extension == ".csv":

            text = _extract_csv(
                file_path
            )

            message = (
                "Data extracted from CSV file."
            )

        elif extension in {
            ".png",
            ".jpg",
            ".jpeg"
        }:

            text, message = _extract_image(
                file_path
            )

        else:

            text = ""

            message = (
                "No extractor is available "
                "for this file type."
            )

        return {
            "supported": True,
            "text": text[:500000],
            "message": message,
            "file_type": file_type,
        }

    except Exception as error:

        return {
            "supported": True,
            "text": "",
            "message": (
                f"Could not read this "
                f"{file_type} file. "
                f"Technical detail: {error}"
            ),
            "file_type": file_type,
        }